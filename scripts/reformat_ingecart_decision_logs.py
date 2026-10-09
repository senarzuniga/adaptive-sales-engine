from __future__ import annotations

import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Literal

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

Language = Literal["es", "en"]
Status = Literal["pending", "in_progress", "confirmed", "blocked", "closed"]
Priority = Literal["critical", "high", "medium", "low", "na"]

ROOT = Path(r"C:\Users\isena\Documents\INGECART\PM")

BLACK = "111111"
ORANGE = "FF6600"
LIGHT_GRAY = "F2F2F2"
DARK_TEXT = "222222"
GRAY_TEXT = "595959"
GREEN = "2E7D32"
BLUE = "1F4E78"
RED = "C00000"
WHITE = "FFFFFF"
BORDER_GRAY = "D0D0D0"


@dataclass
class Point:
    number: int
    title: str
    description: str
    owner: str
    due: str
    status: Status
    priority: Priority
    action: str


@dataclass
class Meta:
    project: str
    customer: str
    owner: str
    revision: str


MOJIBAKE = {
    "Ã¡": "á",
    "Ã©": "é",
    "Ã­": "í",
    "Ã³": "ó",
    "Ãº": "ú",
    "Ã": "Á",
    "Ã‰": "É",
    "Ã": "Í",
    "Ã“": "Ó",
    "Ãš": "Ú",
    "Ã±": "ñ",
    "Ã‘": "Ñ",
    "Ã¼": "ü",
    "Ãœ": "Ü",
    "â€¢": "·",
    "â€“": "–",
    "â€”": "—",
    "â€˜": "‘",
    "â€™": "’",
    "â€œ": "“",
    "â€�": "”",
    "â€¦": "…",
    "Â°": "°",
    "Â·": "·",
    "Â": "",
    "ï»¿": "",
    "\u00a0": " ",
}


def clean(text: object) -> str:
    value = "" if text is None else str(text)
    for src, dst in MOJIBAKE.items():
        value = value.replace(src, dst)
    value = value.replace("�", "·")
    value = re.sub(r"\s+", " ", value).strip()
    return unicodedata.normalize("NFC", value)


def normalize_key(value: str) -> str:
    t = clean(value).upper()
    t = "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn")
    return re.sub(r"[^A-Z0-9]+", " ", t).strip()


def lang_labels(language: Language):
    if language == "es":
        return {
            "sheet": "Registro de decisiones",
            "title": "Registro de Decisiones",
            "subtitle_suffix": "estado y seguimiento",
            "project": "Proyecto",
            "customer": "Cliente",
            "document": "Documento",
            "date": "Fecha",
            "owner": "Responsable",
            "revision": "Rev.",
            "summary": "Resumen",
            "open": "Abiertos",
            "doc_value": "Registro de decisiones del proyecto",
            "footer": "INGECART · Engineering & Auditing — Registro de decisiones del proyecto.",
            "note": "Nota: si no se indica lo contrario, este registro se entrega siempre en formato INGECART oficial.",
            "section_confirmed": "CONFIRMADAS",
            "section_critical": "PENDIENTES · PRIORIDAD CRÍTICA",
            "section_normal": "PENDIENTES · PRIORIDAD ALTA / MEDIA",
            "headers": ["#", "Punto de decisión", "Descripción", "Responsable", "Fecha límite / confirmación", "Estado", "Prioridad", "Acción realizada"],
            "status": {"pending": "PENDIENTE", "in_progress": "EN CURSO", "confirmed": "CONFIRMADO", "blocked": "BLOQUEADO", "closed": "CERRADO"},
            "priority": {"critical": "CRÍTICA", "high": "ALTA", "medium": "MEDIA", "low": "BAJA", "na": "N/A"},
            "summary_words": ("confirmados", "pendientes", "bloqueados", "abiertos", "críticos"),
            "tbd": "Por definir",
        }
    return {
        "sheet": "Decision Log",
        "title": "Decision Log",
        "subtitle_suffix": "status and follow-up",
        "project": "Project",
        "customer": "Customer",
        "document": "Document",
        "date": "Date",
        "owner": "Owner",
        "revision": "Rev.",
        "summary": "Summary",
        "open": "Open",
        "doc_value": "Project decision log",
        "footer": "INGECART · Engineering & Auditing — Project decision log.",
        "note": "Note: unless stated otherwise, this register is always delivered in official INGECART format.",
        "section_confirmed": "CONFIRMED",
        "section_critical": "PENDING · CRITICAL PRIORITY",
        "section_normal": "PENDING · HIGH / MEDIUM PRIORITY",
        "headers": ["#", "Decision point", "Description", "Owner", "Due / confirmation date", "Status", "Priority", "Action taken"],
        "status": {"pending": "PENDING", "in_progress": "IN PROGRESS", "confirmed": "CONFIRMED", "blocked": "BLOCKED", "closed": "CLOSED"},
        "priority": {"critical": "CRITICAL", "high": "HIGH", "medium": "MEDIUM", "low": "LOW", "na": "N/A"},
        "summary_words": ("confirmed", "pending", "blocked", "open", "critical"),
        "tbd": "TBD",
    }


def parse_date_string(value: object, fallback: str) -> str:
    if value is None:
        return fallback
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = clean(value)
    if not text:
        return fallback
    return text


def canonical_status(raw: object) -> Status | None:
    key = normalize_key(str(raw or ""))
    mapping = {
        "PENDIENTE": "pending",
        "PENDING": "pending",
        "EN CURSO": "in_progress",
        "IN PROGRESS": "in_progress",
        "CONFIRMADO": "confirmed",
        "CONFIRMED": "confirmed",
        "CERRADO": "closed",
        "CLOSED": "closed",
        "BLOQUEADO": "blocked",
        "BLOCKED": "blocked",
    }
    return mapping.get(key)


def canonical_priority(raw: object) -> Priority:
    key = normalize_key(str(raw or ""))
    mapping = {
        "CRITICA": "critical",
        "CRITICAL": "critical",
        "ALTA": "high",
        "HIGH": "high",
        "MEDIA": "medium",
        "MEDIUM": "medium",
        "BAJA": "low",
        "LOW": "low",
        "N A": "na",
        "NA": "na",
    }
    return mapping.get(key, "medium")


def extract_meta(ws) -> Meta:
    project = clean(ws["C12"].value) or clean(ws["B3"].value) or "Project"
    customer = clean(ws["F12"].value) or clean(ws["D3"].value) or "—"
    owner = clean(ws["C14"].value) or clean(ws["B5"].value) or "INGECART TEAM"
    revision = clean(ws["F14"].value) or clean(ws["D5"].value) or "1"
    return Meta(project=project, customer=customer, owner=owner, revision=revision)


def is_section_header(value: str) -> bool:
    key = normalize_key(value)
    section_tokens = {
        "CONFIRMADAS",
        "CONFIRMED",
        "PENDIENTES PRIORIDAD CRITICA",
        "PENDING CRITICAL PRIORITY",
        "PENDIENTES PRIORIDAD ALTA MEDIA",
        "PENDING HIGH MEDIUM PRIORITY",
    }
    return key in section_tokens


def extract_points(ws, language: Language) -> list[Point]:
    points: list[Point] = []
    index = 1
    for row in range(1, ws.max_row + 1):
        a = ws.cell(row, 1).value
        b = clean(ws.cell(row, 2).value)
        c = clean(ws.cell(row, 3).value)
        d = clean(ws.cell(row, 4).value)
        e = ws.cell(row, 5).value
        f = ws.cell(row, 6).value
        g = ws.cell(row, 7).value
        h = clean(ws.cell(row, 8).value)

        status = canonical_status(f)
        if not status:
            # fallback for legacy layouts where status might be in column E
            status = canonical_status(e)
            if status:
                e, f, g, h = ws.cell(row, 4).value, ws.cell(row, 5).value, ws.cell(row, 6).value, clean(ws.cell(row, 7).value)
                d = clean(ws.cell(row, 3).value)
                c = clean(ws.cell(row, 2).value)
                b = clean(ws.cell(row, 1).value)

        if not status:
            continue
        if not b or is_section_header(b):
            continue

        points.append(
            Point(
                number=index,
                title=b,
                description=c,
                owner=d or "INGECART TEAM",
                due=parse_date_string(e, "Por definir" if language == "es" else "TBD"),
                status=status,
                priority=canonical_priority(g),
                action=h,
            )
        )
        index += 1
    return points


def section_of(point: Point):
    if point.status in {"confirmed", "closed"}:
        return "confirmed"
    if point.status == "blocked" or point.priority == "critical":
        return "critical"
    return "normal"


def style_row(ws, row: int, blocked: bool):
    fill = PatternFill("solid", fgColor=WHITE if row % 2 == 0 else LIGHT_GRAY)
    border = Border(
        left=Side(style="thin", color=BORDER_GRAY),
        right=Side(style="thin", color=BORDER_GRAY),
        top=Side(style="thin", color=BORDER_GRAY),
        bottom=Side(style="thin", color=BORDER_GRAY),
    )
    for col in "ABCDEFGH":
        c = ws[f"{col}{row}"]
        c.fill = fill
        c.border = border
        c.font = Font(name="Arial", size=10, color=DARK_TEXT)
        c.alignment = Alignment(vertical="top", horizontal="left", wrap_text=True)
    for col in ("A", "E", "F", "G"):
        ws[f"{col}{row}"].alignment = Alignment(vertical="top", horizontal="center", wrap_text=True)
    if blocked:
        ws[f"F{row}"].fill = PatternFill("solid", fgColor=RED)
        ws[f"F{row}"].font = Font(name="Arial", size=10, bold=True, color=WHITE)


def write_ingecart(path: Path, language: Language, meta: Meta, points: list[Point]):
    l = lang_labels(language)
    wb = Workbook()
    ws = wb.active
    ws.title = l["sheet"]
    ws.sheet_view.showGridLines = False

    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 38
    ws.column_dimensions["C"].width = 62
    ws.column_dimensions["D"].width = 20
    ws.column_dimensions["E"].width = 19
    ws.column_dimensions["F"].width = 15
    ws.column_dimensions["G"].width = 13
    ws.column_dimensions["H"].width = 44

    ws.merge_cells("A1:H1")
    ws["A1"].fill = PatternFill("solid", fgColor=BLACK)
    ws.merge_cells("A2:H2")
    ws["A2"].fill = PatternFill("solid", fgColor=ORANGE)
    ws.row_dimensions[2].height = 5

    ws.merge_cells("A4:C4")
    ws["A4"] = "ingecart · Engineering & Auditing"
    ws["A4"].font = Font(name="Arial", size=12, bold=True, color=DARK_TEXT)

    ws.merge_cells("E4:H4")
    ws["E4"] = l["title"]
    ws["E4"].font = Font(name="Arial", size=18, bold=True, color=DARK_TEXT)
    ws["E4"].alignment = Alignment(horizontal="right")

    ws.merge_cells("E5:H5")
    ws["E5"] = f"{meta.customer} · {meta.project} · {l['subtitle_suffix']}"
    ws["E5"].font = Font(name="Arial", size=11, bold=True, color=ORANGE)
    ws["E5"].alignment = Alignment(horizontal="right")

    for r in range(12, 16):
        for c in ("B", "C", "E", "F"):
            cell = ws[f"{c}{r}"]
            cell.fill = PatternFill("solid", fgColor=LIGHT_GRAY)
            cell.font = Font(name="Arial", size=10, color=DARK_TEXT, bold=c in {"B", "E"})
            cell.border = Border(
                left=Side(style="thin", color=ORANGE if c in {"B", "E"} else BORDER_GRAY),
                bottom=Side(style="thin", color=BORDER_GRAY),
            )

    ws["B12"] = l["project"]
    ws["C12"] = meta.project
    ws["E12"] = l["customer"]
    ws["F12"] = meta.customer
    ws["B13"] = l["document"]
    ws["C13"] = l["doc_value"]
    ws["E13"] = l["date"]
    ws["F13"] = date.today()
    ws["F13"].number_format = "DD/MM/YYYY" if language == "es" else "DD-MMM-YYYY"
    ws["B14"] = l["owner"]
    ws["C14"] = meta.owner
    ws["E14"] = l["revision"]
    ws["F14"] = meta.revision

    confirmed = len([p for p in points if p.status in {"confirmed", "closed"}])
    blocked = len([p for p in points if p.status == "blocked"])
    open_count = len([p for p in points if p.status not in {"confirmed", "closed"}])
    critical = len([p for p in points if p.status not in {"confirmed", "closed"} and p.priority == "critical"])
    c1, c2, c3, c4, c5 = l["summary_words"]

    ws["B15"] = l["summary"]
    ws["C15"] = f"{confirmed} {c1} · {open_count} {c2} · {blocked} {c3}"
    ws["E15"] = l["open"]
    ws["F15"] = f"{open_count} {c4} · {critical} {c5}"

    for i, head in enumerate(l["headers"], start=1):
        cell = ws.cell(17, i, head)
        cell.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLACK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[17].height = 22

    grouped = {
        "confirmed": [p for p in points if section_of(p) == "confirmed"],
        "critical": [p for p in points if section_of(p) == "critical"],
        "normal": [p for p in points if section_of(p) == "normal"],
    }
    row = 18
    n = 1
    for key, label in (("confirmed", l["section_confirmed"]), ("critical", l["section_critical"]), ("normal", l["section_normal"])):
        ws.merge_cells(f"A{row}:H{row}")
        ws[f"A{row}"] = label
        ws[f"A{row}"].font = Font(name="Arial", size=10, bold=True, color=WHITE)
        ws[f"A{row}"].fill = PatternFill("solid", fgColor=ORANGE)
        row += 1
        for p in grouped[key]:
            style_row(ws, row, p.status == "blocked")
            ws[f"A{row}"] = n
            ws[f"B{row}"] = p.title
            ws[f"C{row}"] = p.description
            ws[f"D{row}"] = p.owner
            ws[f"E{row}"] = p.due or l["tbd"]
            ws[f"F{row}"] = l["status"][p.status]
            ws[f"G{row}"] = l["priority"][p.priority]
            ws[f"H{row}"] = clean(p.action).upper()
            if p.status in {"confirmed", "closed"}:
                ws[f"F{row}"].font = Font(name="Arial", size=10, bold=True, color=GREEN)
            elif p.status == "in_progress":
                ws[f"F{row}"].font = Font(name="Arial", size=10, bold=True, color=BLUE)
            elif p.status == "pending":
                ws[f"F{row}"].font = Font(name="Arial", size=10, bold=True, color=ORANGE)
            if p.priority == "critical":
                ws[f"G{row}"].font = Font(name="Arial", size=10, bold=True, color=RED)
            elif p.priority == "high":
                ws[f"G{row}"].font = Font(name="Arial", size=10, bold=True, color=ORANGE)
            elif p.priority == "medium":
                ws[f"G{row}"].font = Font(name="Arial", size=10, bold=True, color=GRAY_TEXT)
            row += 1
            n += 1

    reserve_start = row
    reserve_end = row + 29
    for rr in range(reserve_start, reserve_end + 1):
        style_row(ws, rr, False)

    status_values = ",".join(l["status"].values())
    priority_values = ",".join(l["priority"].values())
    dv_status = DataValidation(type="list", formula1=f"\"{status_values}\"", allow_blank=True)
    dv_priority = DataValidation(type="list", formula1=f"\"{priority_values}\"", allow_blank=True)
    ws.add_data_validation(dv_status)
    ws.add_data_validation(dv_priority)
    dv_status.add(f"F18:F{reserve_end}")
    dv_priority.add(f"G18:G{reserve_end}")

    ws.conditional_formatting.add(
        f"F18:F{reserve_end}",
        CellIsRule(operator="equal", formula=[f"\"{l['status']['blocked']}\""], stopIfTrue=True, font=Font(bold=True, color=WHITE), fill=PatternFill("solid", fgColor=RED)),
    )

    footer_row = reserve_end + 2
    ws.merge_cells(f"A{footer_row}:H{footer_row}")
    ws[f"A{footer_row}"] = l["footer"]
    ws[f"A{footer_row}"].font = Font(name="Arial", size=10, italic=True, color=GRAY_TEXT)
    ws[f"A{footer_row}"].border = Border(top=Side(style="thin", color=ORANGE))
    ws.merge_cells(f"A{footer_row + 1}:H{footer_row + 1}")
    ws[f"A{footer_row + 1}"] = l["note"]
    ws[f"A{footer_row + 1}"].font = Font(name="Arial", size=10, italic=True, color=GRAY_TEXT)

    ws.freeze_panes = "A18"
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "17:17"
    wb.save(path)


def detect_language(path: Path) -> Language:
    upper = path.name.upper()
    if upper.endswith("_INGECART_EN.XLSX"):
        return "en"
    return "es"


def is_candidate(path: Path) -> bool:
    return bool(re.match(r"^(Registro_Decisiones_.*_INGECART_ES\.xlsx|Decision_Log_.*_INGECART_EN\.xlsx)$", path.name))


def main():
    files = sorted(p for p in ROOT.rglob("*.xlsx") if is_candidate(p) and "IP SR1400" not in str(p))
    reformatted = 0
    for file in files:
        language = detect_language(file)
        wb = load_workbook(file, data_only=False)
        ws = wb.active
        meta = extract_meta(ws)
        points = extract_points(ws, language)
        if not points:
            print(f"SKIP (no points parsed): {file}")
            continue
        write_ingecart(file, language, meta, points)
        reformatted += 1
        print(f"REFORMATTED: {file}")

    # remove duplicate legacy valley naming variants if present
    duplicates = [
        ROOT / "CARTONAJES FONT" / "LINEA POTENCIA VAHLEY" / "Decision_Log_Cartonajes_Font_Linea_Potencia_Valley_INGECART_EN.xlsx",
        ROOT / "CARTONAJES FONT" / "LINEA POTENCIA VAHLEY" / "Registro_Decisiones_Cartonajes_Font_Linea_Potencia_Valley_INGECART_ES.xlsx",
    ]
    for dup in duplicates:
        if dup.exists():
            dup.unlink()
            print(f"REMOVED DUPLICATE: {dup}")

    print(f"TOTAL REFORMATTED: {reformatted}")


if __name__ == "__main__":
    main()
